from openpyxl import load_workbook

wb = load_workbook('Amazon_EC2_Instances_BulkUpload_Template_Commercial.xlsx')
ws = wb['Inputs']

print('Checking all rows with data in the template:')
print(f'Max row with data: {ws.max_row}')
print(f'Max column with data: {ws.max_column}')

print('\nColumn 9 (Assumed Usage) values:')
for row in range(1, ws.max_row + 1):
    cell = ws.cell(row=row, column=9)
    if cell.value is not None:
        print(f'Row {row}: value={cell.value}')
